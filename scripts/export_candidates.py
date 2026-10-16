#!/usr/bin/env python3
"""
IEEE-VIT Enrollments - Candidate Data Exporter
Exports candidate personal data, selected domains, candidate submitted projects,
assigned recruitment tasks, application answers, slot bookings, and interview reviews
into styled Excel (.xlsx) or CSV (.csv) format.

Usage:
    python scripts/export_candidates.py [options]

Examples:
    python scripts/export_candidates.py                       # Export to timestamped Excel file
    python scripts/export_candidates.py --format csv          # Export to CSV
    python scripts/export_candidates.py --domain "Web Development"
    python scripts/export_candidates.py --stage task
    python scripts/export_candidates.py --output candidates.xlsx
    python scripts/export_candidates.py --demo                # Generate demo sample without live DB
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
except ImportError:
    psycopg2 = None

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None


# Known application questions mapping for friendly column headers
QUESTION_HEADER_MAP = {
    "whyIeee": "Why IEEE-VIT?",
    "currentSkills": "Current Skills & Tools",
    "exploreNext": "Skill to Explore Next",
    "pastExperience": "Past Experience",
    "commitment": "Weekly Commitment",
}


def sanitize_formula_injection(value: Any) -> Any:
    """Neutralize spreadsheet formula injection characters (=, +, -, @)."""
    if isinstance(value, str) and value:
        if value[0] in ("=", "+", "-", "@"):
            return f"'{value}"
    return value


def format_dt(dt: Optional[datetime]) -> str:
    """Format datetime into standard readable string."""
    if not dt:
        return ""
    if isinstance(dt, str):
        return dt
    # Format as YYYY-MM-DD HH:MM UTC
    return dt.strftime("%Y-%m-%d %H:%M")


def compute_stage(app_submitted: bool, task_submitted: bool, slot_id: Optional[str]) -> str:
    """
    Derives candidate stage matching backend candidate.service.js:
      1. application: form not submitted
      2. task: form submitted, task not submitted
      3. slot: task submitted, slot not booked
      4. interview: slot booked
    """
    if not app_submitted:
        return "application"
    if not task_submitted:
        return "task"
    if not slot_id:
        return "slot"
    return "interview"


def compute_task_status(task_title: Optional[str], submitted_at: Optional[datetime], deadline: Optional[datetime]) -> str:
    """
    Derives task status matching backend candidate.service.js:
      - awaiting: applied, but no task assigned yet
      - submitted: task submitted
      - expired: deadline passed without submission
      - active: task assigned and within deadline
    """
    if not task_title:
        return "awaiting"
    if submitted_at:
        return "submitted"
    if deadline:
        now = datetime.now(timezone.utc)
        if deadline.tzinfo is None:
            deadline = deadline.replace(tzinfo=timezone.utc)
        if now > deadline:
            return "expired"
    return "active"


def resolve_database_url(cli_db_url: Optional[str], cli_env_file: Optional[str]) -> str:
    """Finds DATABASE_URL from CLI, specified env file, or standard paths."""
    if cli_db_url:
        return cli_db_url

    # Check CLI env file first
    if cli_env_file and Path(cli_env_file).exists() and load_dotenv:
        load_dotenv(cli_env_file)

    # Search standard paths for .env
    current_dir = Path(__file__).resolve().parent
    repo_root = current_dir.parent
    backend_env = repo_root / ".env"
    local_env = current_dir / ".env"
    parent_env = repo_root.parent / ".env"

    for candidate_env in [backend_env, local_env, parent_env]:
        if candidate_env.exists() and load_dotenv:
            load_dotenv(candidate_env)
            break

    db_url = os.getenv("DATABASE_URL", "").strip()
    return db_url


def fetch_candidate_data(db_url: str, domain_filter: Optional[str], stage_filter: Optional[str], include_incomplete: bool) -> List[Dict[str, Any]]:
    """Connects to PostgreSQL and extracts complete candidate details."""
    if not psycopg2:
        raise RuntimeError("Missing 'psycopg2' package. Install with: pip install psycopg2-binary")

    if not db_url or db_url.startswith("pglite:"):
        raise ValueError(
            "Valid Postgres DATABASE_URL not found.\n"
            "Please provide a PostgreSQL connection URL via:\n"
            "  --db-url 'postgresql://user:pass@host/db'\n"
            "  or set DATABASE_URL in .env\n"
            "  (Or use --demo to test with sample data)"
        )

    # Clean connection string parameters if needed for psycopg2
    # e.g., psycopg2 doesn't recognize channel_binding
    clean_url = db_url.replace("channel_binding=require&", "").replace("&channel_binding=require", "").replace("channel_binding=require", "")

    query = """
    SELECT
        u.id AS user_id,
        u.name,
        u.reg_no,
        u.email,
        u.created_at AS user_created_at,
        u.slot_id,
        a.phone,
        a.github_profile,
        a.linkedin_profile,
        a.project_link AS candidate_project_link,
        a.project_description AS candidate_project_description,
        a.first_domain,
        a.second_domain,
        a.answers,
        a.management_question_id,
        a.management_question,
        a.management_answer,
        a.submitted_at AS application_submitted_at,
        ta.title AS task_title,
        ta.description AS task_description,
        ta.domain AS task_domain,
        ta.assigned_at AS task_assigned_at,
        ta.deadline AS task_deadline,
        ta.submission_link AS task_submission_link,
        ta.submission_notes AS task_submission_notes,
        ta.submitted_at AS task_submitted_at,
        s.start_time AS slot_start_time,
        s.duration_minutes AS slot_duration_minutes,
        s.meet_link AS slot_meet_link,
        ir.status AS interview_status,
        ir.rating AS interview_rating,
        ir.remarks AS interview_remarks,
        ir.reviewed_at AS interview_reviewed_at
    FROM users u
    LEFT JOIN applications a ON a.user_id = u.id
    LEFT JOIN task_assignments ta ON ta.user_id = u.id
    LEFT JOIN slots s ON s.id = u.slot_id
    LEFT JOIN interview_reviews ir ON ir.user_id = u.id
    WHERE 1=1
    """

    params: List[Any] = []

    if not include_incomplete:
        query += " AND a.submitted_at IS NOT NULL"

    if domain_filter:
        query += " AND (a.first_domain ILIKE %s OR a.second_domain ILIKE %s)"
        params.extend([f"%{domain_filter}%", f"%{domain_filter}%"])

    query += " ORDER BY u.created_at ASC"

    conn = psycopg2.connect(clean_url)
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query, params)
            rows = [dict(r) for r in cur.fetchall()]
    finally:
        conn.close()

    # Post-process rows: derive stages & task statuses
    processed = []
    for r in rows:
        app_sub = r.get("application_submitted_at") is not None
        task_sub = r.get("task_submitted_at") is not None
        stage = compute_stage(app_sub, task_sub, r.get("slot_id"))
        task_status = compute_task_status(r.get("task_title"), r.get("task_submitted_at"), r.get("task_deadline"))

        r["derived_stage"] = stage
        r["derived_task_status"] = task_status

        # If filtered by stage, skip non-matching
        if stage_filter and stage.lower() != stage_filter.lower():
            continue

        processed.append(r)

    return processed


def get_demo_candidates() -> List[Dict[str, Any]]:
    """Provides realistic mock candidate data for verification and demonstrations."""
    now = datetime.now(timezone.utc)
    return [
        {
            "user_id": "11111111-1111-1111-1111-111111111111",
            "name": "Aarav Sharma",
            "reg_no": "24BCE1042",
            "email": "aarav.sharma2024@vitstudent.ac.in",
            "user_created_at": now,
            "slot_id": "slot-uuid-1",
            "phone": "+91 98765 43210",
            "github_profile": "https://github.com/aarav-sharma",
            "linkedin_profile": "https://linkedin.com/in/aarav-sharma",
            "candidate_project_link": "https://github.com/aarav-sharma/campus-events-hub",
            "candidate_project_description": "Full-stack React & Node.js application to register and browse student technical events.",
            "first_domain": "Web Development",
            "second_domain": "App Development",
            "answers": [
                {"questionId": "whyIeee", "answer": "I want to collaborate on high-impact technical projects and learn modern system design."},
                {"questionId": "currentSkills", "answer": "JavaScript, TypeScript, React, Express, PostgreSQL, Git."},
                {"questionId": "exploreNext", "answer": "Distributed systems, Docker, and Kubernetes deployment pipelines."},
                {"questionId": "pastExperience", "answer": "Built the frontend for our school's annual science fest portal."},
                {"questionId": "commitment", "answer": "12-15 hours/week."},
            ],
            "management_question": "Two key members of your team refuse to collaborate because of personal differences. How would you handle this?",
            "management_answer": "I would hold a 1-on-1 discussion with each to understand their concerns, set neutral objective boundaries, and focus on team deadlines.",
            "application_submitted_at": now,
            "task_title": "Build a Real-Time Collaborative Canvas",
            "task_description": "Create a multi-user whiteboard using WebSockets and Canvas API with brush color and undo/redo.",
            "task_domain": "Web Development",
            "task_assigned_at": now,
            "task_deadline": now,
            "task_submission_link": "https://github.com/aarav-sharma/ieee-canvas-task",
            "task_submission_notes": "Implemented room support and custom brush sizes using Socket.IO.",
            "task_submitted_at": now,
            "slot_start_time": now,
            "slot_duration_minutes": 20,
            "slot_meet_link": "https://meet.google.com/abc-defg-hij",
            "interview_status": "selected",
            "interview_rating": 5,
            "interview_remarks": "Strong fundamentals in web security and WebSocket state synchronization. Excellent cultural fit.",
            "interview_reviewed_at": now,
            "derived_stage": "interview",
            "derived_task_status": "submitted",
        },
        {
            "user_id": "22222222-2222-2222-2222-222222222222",
            "name": "Diya Patel",
            "reg_no": "24BCI0219",
            "email": "diya.patel2024@vitstudent.ac.in",
            "user_created_at": now,
            "slot_id": None,
            "phone": "+91 91234 56789",
            "github_profile": "https://github.com/diyapatel",
            "linkedin_profile": "https://linkedin.com/in/diya-patel",
            "candidate_project_link": "https://github.com/diyapatel/retinal-disease-classifier",
            "candidate_project_description": "Deep learning ResNet-50 model trained on fundus images to classify diabetic retinopathy stages.",
            "first_domain": "Machine Learning",
            "second_domain": "Cyber Security",
            "answers": [
                {"questionId": "whyIeee", "answer": "To work with research-oriented developers and publish innovative IEEE conference papers."},
                {"questionId": "currentSkills", "answer": "Python, PyTorch, Scikit-learn, OpenCV, Pandas, NumPy."},
                {"questionId": "exploreNext", "answer": "Vision Transformers (ViT) and Multimodal LLM fine-tuning."},
                {"questionId": "pastExperience", "answer": "Top 10 in college Smart Healthcare Hackathon 2024."},
                {"questionId": "commitment", "answer": "10 hours/week."},
            ],
            "management_question": "During a live online contest, the platform crashes. How would you handle this?",
            "management_answer": "Post an immediate status banner, pause the contest clock, address inquiries transparently, and extend the timer upon recovery.",
            "application_submitted_at": now,
            "task_title": "Audio Deepfake Detection Pipeline",
            "task_description": "Develop a classifier to distinguish synthetic voice generation from human speech using spectrogram features.",
            "task_domain": "Machine Learning",
            "task_assigned_at": now,
            "task_deadline": now,
            "task_submission_link": "https://github.com/diyapatel/audio-deepfake-detect",
            "task_submission_notes": "Extracted MFCCs and mel-spectrograms, evaluated on ASVspoof dataset with 92% ROC-AUC.",
            "task_submitted_at": now,
            "slot_start_time": None,
            "slot_duration_minutes": None,
            "slot_meet_link": None,
            "interview_status": "pending",
            "interview_rating": None,
            "interview_remarks": None,
            "interview_reviewed_at": None,
            "derived_stage": "slot",
            "derived_task_status": "submitted",
        },
        {
            "user_id": "33333333-3333-3333-3333-333333333333",
            "name": "Rohan Iyer",
            "reg_no": "24BEC0881",
            "email": "rohan.iyer2024@vitstudent.ac.in",
            "user_created_at": now,
            "slot_id": None,
            "phone": "+91 99887 76655",
            "github_profile": "https://github.com/rohaniyer",
            "linkedin_profile": "https://linkedin.com/in/rohan-iyer",
            "candidate_project_link": "https://github.com/rohaniyer/smart-irrigation-node",
            "candidate_project_description": "ESP32 based solar-powered agricultural sensor transmitting soil moisture and temperature via LoRa.",
            "first_domain": "Electronics",
            "second_domain": "App Development",
            "answers": [
                {"questionId": "whyIeee", "answer": "IEEE has the strongest hardware and embedded development lab on campus."},
                {"questionId": "currentSkills", "answer": "Embedded C, Arduino, KiCad, ESP32, STM32, MQTT."},
                {"questionId": "exploreNext", "answer": "FPGA programming with Verilog and TinyML on microcontrollers."},
                {"questionId": "pastExperience", "answer": "Designed custom dual-layer PCB for college robotics team line follower."},
                {"questionId": "commitment", "answer": "14 hours/week."},
            ],
            "management_question": "Two days before an event, main sponsor pulls out. What do you do in 48 hours?",
            "management_answer": "Prioritize vital event operations, identify alternative local sponsors and university student activity grants, and adjust collateral.",
            "application_submitted_at": now,
            "task_title": "Autonomous Indoor Air Quality Monitor",
            "task_description": "Design an ESP32 schematic and firmware to read MQ-135 and BME280 sensors, publishing metrics to MQTT dashboard.",
            "task_domain": "Electronics",
            "task_assigned_at": now,
            "task_deadline": now,
            "task_submission_link": None,
            "task_submission_notes": None,
            "task_submitted_at": None,
            "slot_start_time": None,
            "slot_duration_minutes": None,
            "slot_meet_link": None,
            "interview_status": "pending",
            "interview_rating": None,
            "interview_remarks": None,
            "interview_reviewed_at": None,
            "derived_stage": "task",
            "derived_task_status": "active",
        },
    ]


def prepare_rows_and_headers(candidates: List[Dict[str, Any]]) -> Tuple[List[str], List[List[Any]]]:
    """Flattens candidate records into tabular headers and rows."""
    # Find all distinct questionIds across candidates
    question_ids = []
    seen_q_ids = set()
    for c in candidates:
        answers = c.get("answers") or []
        if isinstance(answers, str):
            try:
                answers = json.loads(answers)
            except Exception:
                answers = []
        for a in answers:
            qid = a.get("questionId")
            if qid and qid not in seen_q_ids:
                seen_q_ids.add(qid)
                question_ids.append(qid)

    # Core Headers
    headers = [
        "Candidate Name",
        "Reg No",
        "Email",
        "Phone",
        "Stage",
        "First Domain",
        "Second Domain",
        "GitHub Profile",
        "LinkedIn Profile",
        # Candidate's personal submitted project
        "Candidate Project Link",
        "Candidate Project Description",
        # Assigned recruitment task
        "Assigned Task Title",
        "Assigned Task Domain",
        "Task Description",
        "Task Assigned At",
        "Task Deadline",
        "Task Status",
        "Task Submission Link",
        "Task Submission Notes",
        "Task Submitted At",
        # Interview details
        "Interview Slot Time",
        "Slot Duration (min)",
        "Slot Meeting Link",
        "Interview Status",
        "Interview Rating",
        "Interview Remarks",
        # Situational Management & Domain Q&A
        "Management Question",
        "Management Answer",
    ]

    for qid in question_ids:
        headers.append(f"Q: {QUESTION_HEADER_MAP.get(qid, qid)}")

    headers.extend(["Application Submitted At", "Account Registered At"])

    rows = []
    for c in candidates:
        answers_list = c.get("answers") or []
        if isinstance(answers_list, str):
            try:
                answers_list = json.loads(answers_list)
            except Exception:
                answers_list = []
        answer_dict = {a.get("questionId"): a.get("answer", "") for a in answers_list if isinstance(a, dict)}

        row = [
            c.get("name", ""),
            c.get("reg_no", "") or "",
            c.get("email", ""),
            c.get("phone", "") or "",
            c.get("derived_stage", ""),
            c.get("first_domain", "") or "",
            c.get("second_domain", "") or "",
            c.get("github_profile", "") or "",
            c.get("linkedin_profile", "") or "",
            # Candidate project
            c.get("candidate_project_link", "") or "",
            c.get("candidate_project_description", "") or "",
            # Assigned task
            c.get("task_title", "") or "",
            c.get("task_domain", "") or "",
            c.get("task_description", "") or "",
            format_dt(c.get("task_assigned_at")),
            format_dt(c.get("task_deadline")),
            c.get("derived_task_status", ""),
            c.get("task_submission_link", "") or "",
            c.get("task_submission_notes", "") or "",
            format_dt(c.get("task_submitted_at")),
            # Interview
            format_dt(c.get("slot_start_time")),
            c.get("slot_duration_minutes", "") or "",
            c.get("slot_meet_link", "") or "",
            c.get("interview_status", "") or "",
            c.get("interview_rating", "") if c.get("interview_rating") is not None else "",
            c.get("interview_remarks", "") or "",
            # Management Q&A
            c.get("management_question", "") or "",
            c.get("management_answer", "") or "",
        ]

        # Specific question answers
        for qid in question_ids:
            row.append(answer_dict.get(qid, ""))

        # Timestamps
        row.append(format_dt(c.get("application_submitted_at")))
        row.append(format_dt(c.get("user_created_at")))

        # Sanitize against CSV formula injection
        row = [sanitize_formula_injection(val) for val in row]
        rows.append(row)

    return headers, rows


def export_to_csv(headers: List[str], rows: List[List[Any]], output_path: Path):
    """Writes candidate data to CSV with UTF-8 BOM encoding."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, mode="w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def export_to_xlsx(headers: List[str], rows: List[List[Any]], output_path: Path):
    """Writes candidate data to a styled Excel workbook."""
    if not openpyxl:
        raise RuntimeError("Missing 'openpyxl' package. Install with: pip install openpyxl")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Participants"

    # Styling elements
    # Dark Navy Header
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    
    # Alternating row fill
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_alignment = Alignment(vertical="top", wrap_text=True)

    # Write Headers
    ws.append(headers)
    ws.row_dimensions[1].height = 32

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write Data Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 24
        is_even = (row_idx % 2 == 0)
        current_fill = zebra_fill if is_even else white_fill

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = current_fill
            cell.border = thin_border
            cell.alignment = data_alignment

    # Auto-adjust column widths with reasonable boundaries
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            # Measure line lengths (if multiline)
            lines = val_str.split("\n")
            longest_line = max(len(l) for l in lines) if lines else len(val_str)
            if longest_line > max_len:
                max_len = longest_line
        
        # Clamp column width between 14 and 45
        adjusted_width = max(min(max_len + 3, 45), 14)
        ws.column_dimensions[col_letter].width = adjusted_width

    # Freeze top row
    ws.freeze_panes = "A2"

    # Add auto-filter
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="IEEE-VIT Enrollments - Export participant records, domains, projects & tasks to Excel or CSV."
    )
    parser.add_argument(
        "--format",
        choices=["xlsx", "csv"],
        default="xlsx",
        help="Export file format: 'xlsx' (default) or 'csv'",
    )
    parser.add_argument(
        "-o", "--output",
        help="Custom output file path (defaults to ./exports/enrollments_YYYY-MM-DD_HHMMSS.<ext>)",
    )
    parser.add_argument(
        "--domain",
        help="Filter candidates by first or second domain (e.g. 'Web Development', 'Design')",
    )
    parser.add_argument(
        "--stage",
        choices=["application", "task", "slot", "interview"],
        help="Filter candidates by recruitment stage",
    )
    parser.add_argument(
        "--include-incomplete",
        action="store_true",
        help="Include users who logged in but haven't submitted the application form",
    )
    parser.add_argument(
        "--db-url",
        help="PostgreSQL connection string override (e.g. postgresql://user:pass@host/db)",
    )
    parser.add_argument(
        "--env-file",
        help="Path to custom .env file to load DATABASE_URL from",
    )
    parser.add_argument(
        "--demo",
        action="store_true",
        help="Run using demo participant data to preview generated spreadsheet without connecting to live database",
    )
    parser.add_argument(
        "-q", "--quiet",
        action="store_true",
        help="Suppress informational stdout output",
    )

    args = parser.parse_args()

    # Determine output file path
    now_str = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    if args.output:
        out_path = Path(args.output).resolve()
    else:
        out_path = Path(f"./exports/enrollments_{now_str}.{args.format}").resolve()

    if not args.quiet:
        print("=" * 65)
        print("  IEEE-VIT Enrollments - Candidate Data Exporter")
        print("=" * 65)
        print(f"[*] Target Format: {args.format.upper()}")
        print(f"[*] Output Path:   {out_path}")
        if args.domain:
            print(f"[*] Domain Filter: {args.domain}")
        if args.stage:
            print(f"[*] Stage Filter:  {args.stage}")
        if args.include_incomplete:
            print("[*] Incomplete:    Including unsubmitted signups")

    # Fetch Data
    if args.demo:
        if not args.quiet:
            print("[*] Mode:          DEMO (using realistic sample data)")
        candidates = get_demo_candidates()
        # Apply filters to demo data if specified
        if args.domain:
            candidates = [
                c for c in candidates
                if args.domain.lower() in (c.get("first_domain", "") or "").lower()
                or args.domain.lower() in (c.get("second_domain", "") or "").lower()
            ]
        if args.stage:
            candidates = [c for c in candidates if c.get("derived_stage") == args.stage]
    else:
        db_url = resolve_database_url(args.db_url, args.env_file)
        if not db_url:
            print("\n[ERROR] No DATABASE_URL found.")
            print("Please provide connection details via:")
            print("  1. --db-url 'postgresql://...'")
            print("  2. Setting DATABASE_URL in .env")
            print("  3. Or test with --demo to preview the export structure.")
            sys.exit(1)

        if not args.quiet:
            # Mask credentials when printing
            display_host = db_url.split("@")[-1] if "@" in db_url else "database"
            print(f"[*] DB Target:     {display_host}")
            print("[*] Fetching candidates from database...")

        try:
            candidates = fetch_candidate_data(
                db_url=db_url,
                domain_filter=args.domain,
                stage_filter=args.stage,
                include_incomplete=args.include_incomplete,
            )
        except Exception as e:
            print(f"\n[ERROR] Failed to query candidate data: {e}")
            sys.exit(1)

    if not candidates:
        print("[!] No candidates found matching the criteria.")
        sys.exit(0)

    if not args.quiet:
        print(f"[*] Processing {len(candidates)} candidate record(s)...")

    headers, rows = prepare_rows_and_headers(candidates)

    try:
        if args.format == "csv":
            export_to_csv(headers, rows, out_path)
        else:
            export_to_xlsx(headers, rows, out_path)
    except Exception as e:
        print(f"\n[ERROR] Failed to write {args.format.upper()} file: {e}")
        sys.exit(1)

    if not args.quiet:
        print(f"[OK] Successfully exported {len(rows)} participant(s) to:")
        print(f"    {out_path}")
        print("=" * 65)


if __name__ == "__main__":
    main()
